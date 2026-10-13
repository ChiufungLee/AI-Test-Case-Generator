"""测试任务「需求原文」上传需求文档：纯文本提取接口、来源文件名落库与列迁移回归。

提取接口不落盘、不落库、不进向量库——文档正文由前端回填到需求原文文本框，
用户确认后再走原有的 JSON 创建接口（契约不变）。
"""

import io

from sqlalchemy import inspect, text

from models.database import _ensure_schema_updates, get_engine, init_db
from services import knowledge_service


def _docx_bytes(paragraphs=None, table=None) -> bytes:
    """用 python-docx 现场生成一个 .docx（走真实解析路径，不打桩）"""
    from docx import Document

    document = Document()
    for paragraph in paragraphs or []:
        document.add_paragraph(paragraph)
    if table:
        rows, columns = len(table), len(table[0])
        docx_table = document.add_table(rows=rows, cols=columns)
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                docx_table.cell(row_index, column_index).text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(client, filename: str, content: bytes, content_type: str = "application/octet-stream"):
    return client.post(
        "/api/workflows/requirement-extract",
        files={"file": (filename, content, content_type)},
    )


def test_extract_requirement_requires_login(client):
    response = _upload(client, "req.txt", "需求：新增登录".encode("utf-8"))
    assert response.status_code == 401


def test_workflow_page_renders_requirement_upload_controls(logged_in_client):
    """新建面板需带上上传按钮、文件输入与解析状态位（模板回归）"""
    response = logged_in_client.get("/workflows")

    assert response.status_code == 200
    assert 'id="requirementUploadBtn"' in response.text
    assert 'id="requirementFileInput"' in response.text
    assert 'accept=".pdf,.docx,.txt,.md"' in response.text
    assert 'id="requirementUploadStatus"' in response.text
    assert 'id="requirementSource"' in response.text


def test_extract_requirement_from_txt(logged_in_client, test_env):
    response = _upload(logged_in_client, "登录需求.txt", "需求：新增手机号验证码登录".encode("utf-8"))

    assert response.status_code == 200
    body = response.json()
    assert body["filename"] == "登录需求.txt"
    assert body["text"] == "需求：新增手机号验证码登录"
    assert body["chars"] == len(body["text"])
    assert body["truncated"] is False

    # 临时文件解析后立即删除，不留孤儿文件
    assert list(test_env["temp_upload_dir"].glob("*")) == []


def test_extract_requirement_decodes_gb18030_txt(logged_in_client):
    """中文需求文档常见 GB18030 编码：不能因 UTF-8 解码失败而报错或乱码"""
    response = _upload(logged_in_client, "需求.txt", "需求：锁定30分钟".encode("gb18030"))

    assert response.status_code == 200
    assert response.json()["text"] == "需求：锁定30分钟"


def test_extract_requirement_normalizes_kangxi_radicals(logged_in_client):
    """PDF/Word 转换常见的康熙部首与兼容汉字：归一为常规汉字（并补简体字形），不计入剔除数

    视觉上与常规汉字几乎一致、码位不同（租⼾ U+2F3E / 接收⼈ U+2F08 / 豈 U+F900），
    不归一时在缺字形的字体里就是页面上看到的方框。
    """
    content = "租\u2f3e、接收\u2f08、\uf900".encode("utf-8")

    response = _upload(logged_in_client, "req.txt", content)

    assert response.status_code == 200
    body = response.json()
    assert body["text"] == "租户、接收人、\u8c48"
    assert body["dropped_chars"] == 0


def test_extract_requirement_drops_unrenderable_symbols(logged_in_client):
    """替换字符/私用区/控制符/零宽符在页面上只会显示成方框：剔除并回显数量"""
    content = "名称\ufffd\u0007\ue000\u200b描述\u0085租户".encode("utf-8")

    response = _upload(logged_in_client, "req.txt", content)

    assert response.status_code == 200
    body = response.json()
    # U+0085(NEL) 是换行语义，转成换行而不算剔除
    assert body["text"] == "名称描述\n租户"
    assert body["dropped_chars"] == 4
    assert body["truncated"] is False


def test_extract_requirement_normalizes_radical_supplement_and_variants(logged_in_client):
    """CJK 部首补充（U+2E80–2EFF 无 NFKC 分解）与部首归一得到的繁体字形也要还原成简体"""
    content = "\u2edc\u4e66\u3001\u90e8\u2ed4\u3001\u79df\u2f3e\u3001\u2f08\u5458".encode("utf-8")

    response = _upload(logged_in_client, "req.txt", content)

    assert response.status_code == 200
    assert response.json()["text"] == "飞书、部门、租户、人员"


def test_extract_requirement_keeps_genuine_traditional_text(logged_in_client):
    """归一化只作用于部首类码位：文档里本来就有的繁体正文必须原样保留"""
    text = "繁體中文與繁體欄位不應被改寫"

    response = _upload(logged_in_client, "req.txt", text.encode("utf-8"))

    assert response.json()["text"] == text


def test_extract_requirement_keeps_fullwidth_punctuation(logged_in_client):
    """归一化只针对兼容汉字：全角标点/数字必须原样保留（不整体 NFKC）"""
    text = "需求：连续５次失败，锁定30分钟。（见附录①）"

    response = _upload(logged_in_client, "req.txt", text.encode("utf-8"))

    assert response.json()["text"] == text


def test_extract_requirement_from_markdown_normalizes_and_caps(logged_in_client, monkeypatch):
    """连续空行被压缩；超过上限截断并回报 truncated"""
    monkeypatch.setattr(knowledge_service, "MAX_REQUIREMENT_DOC_CHARS", 20)
    content = b"# \xe9\x9c\x80\xe6\xb1\x82\n\n\n\n\xe6\xad\xa3\xe6\x96\x87 " + b"a" * 40

    response = _upload(logged_in_client, "req.md", content)

    assert response.status_code == 200
    body = response.json()
    assert body["truncated"] is True
    assert len(body["text"]) == 20
    assert "\n\n\n" not in body["text"]


def test_extract_requirement_from_docx_includes_tables(logged_in_client):
    """Word 需求文档：段落与表格都要进正文，表格渲染为 Markdown（业务规则常写在表格里）"""
    content = _docx_bytes(
        paragraphs=["需求：新增验证码登录", "", "业务规则："],
        table=[["规则", "说明"], ["锁定", "连续5次错误锁定30分钟"]],
    )

    response = _upload(
        logged_in_client,
        "需求.docx",
        content,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert response.status_code == 200
    extracted = response.json()["text"]
    assert "需求：新增验证码登录" in extracted
    assert "业务规则：" in extracted
    assert "| 规则 | 说明 |" in extracted
    assert "| --- | --- |" in extracted
    assert "| 锁定 | 连续5次错误锁定30分钟 |" in extracted


def test_extract_requirement_pdf_table_kept_as_markdown(logged_in_client, tmp_path, make_layout_pdf):
    """PDF 按列写入内容流的表格：回填文本要还原成 Markdown 表格，水印不进正文"""
    pdf_path = make_layout_pdf(
        tmp_path / "配置-编辑.pdf",
        rows=[["name", "desc", "note"], ["msg type", "same as create"], ["remark", "same as create"]],
    )
    with open(pdf_path, "rb") as handle:
        content = handle.read()

    response = _upload(logged_in_client, "配置-编辑.pdf", content, "application/pdf")

    assert response.status_code == 200
    body = response.json()
    lines = body["text"].splitlines()
    assert lines[0] == "| name | desc | note |"
    assert lines[1] == "| --- | --- | --- |"
    assert "| msg type | same as create |  |" in lines
    assert "CONFIDENTIAL" not in body["text"]


def test_extract_requirement_from_pdf_uses_document_processor(logged_in_client, document_processor):
    """PDF 走 DocumentProcessor 三级加载（conftest 已将 load_pdf 打桩）"""
    response = _upload(logged_in_client, "req.pdf", b"%PDF-1.4\n%stub\n", "application/pdf")

    assert response.status_code == 200
    assert response.json()["text"] == "pdf content"


def test_extract_requirement_rejects_unsupported_extension(logged_in_client):
    response = _upload(logged_in_client, "需求.doc", b"legacy binary")

    assert response.status_code == 400
    assert "仅支持上传" in response.json()["error"]


def test_extract_requirement_rejects_empty_file(logged_in_client):
    response = _upload(logged_in_client, "req.txt", b"")

    assert response.status_code == 400
    assert response.json()["error"] == "文件内容不能为空"


def test_extract_requirement_rejects_unparsable_document(
    logged_in_client, document_processor, monkeypatch
):
    monkeypatch.setattr("utils.file_handle.DocumentProcessor.load_pdf", lambda self, file_path: [])

    response = _upload(logged_in_client, "scan.pdf", b"%PDF-1.4\n%stub\n", "application/pdf")

    assert response.status_code == 400
    assert response.json()["error"] == "无法从文档中提取到文本内容"


def test_extract_requirement_maps_parser_crash_to_400(
    logged_in_client, document_processor, monkeypatch, test_env
):
    """损坏/加密文档：解析异常应转成可行动的 400，而不是无信息的 500，且清理临时文件"""

    def _boom(self, file_path):
        raise RuntimeError("corrupted pdf")

    monkeypatch.setattr("utils.file_handle.DocumentProcessor.load_pdf", _boom)

    response = _upload(logged_in_client, "broken.pdf", b"%PDF-1.4\nbroken\n", "application/pdf")

    assert response.status_code == 400
    assert "文档解析失败" in response.json()["error"]
    assert list(test_env["temp_upload_dir"].glob("*")) == []


def test_create_workflow_records_requirement_file_name(logged_in_client):
    response = logged_in_client.post(
        "/api/workflows",
        json={
            "name": "登录需求",
            "requirement_text": "需求：新增手机号验证码登录",
            "knowledge_base_id": None,
            "requirement_file_name": "登录需求.docx",
        },
    )

    assert response.status_code == 200
    workflow = response.json()
    assert workflow["requirement_file_name"] == "登录需求.docx"

    detail = logged_in_client.get(f"/api/workflows/{workflow['id']}").json()
    assert detail["requirement_file_name"] == "登录需求.docx"

    listed = logged_in_client.get("/api/workflows").json()["workflows"]
    assert listed[0]["requirement_file_name"] == "登录需求.docx"


def test_create_workflow_without_document_has_null_file_name(logged_in_client):
    response = logged_in_client.post(
        "/api/workflows",
        json={"name": "粘贴需求", "requirement_text": "需求：新增登录", "knowledge_base_id": None},
    )

    assert response.status_code == 200
    assert response.json()["requirement_file_name"] is None


def test_create_workflow_rejects_overlong_file_name(logged_in_client):
    """来源文件名只是溯源标签，超长直接 422 拒绝而不是写坏列"""
    response = logged_in_client.post(
        "/api/workflows",
        json={
            "name": "x",
            "requirement_text": "需求：新增登录",
            "requirement_file_name": "名" * 256,
        },
    )

    assert response.status_code == 422


def test_schema_update_adds_requirement_file_name_column(db_session, make_user):
    """模拟旧版库：删列后重跑迁移 → 列补回，存量任务保持 NULL"""
    alice = make_user("legacy_wf_owner", "secret123")
    db_session.execute(text("ALTER TABLE workflows DROP COLUMN requirement_file_name"))
    db_session.commit()
    db_session.execute(
        text(
            "INSERT INTO workflows (id, user_id, name, requirement_text, status, created_at, updated_at) "
            "VALUES ('legacy-workflow', :user_id, '旧任务', '需求：历史粘贴', 'created', "
            "CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)"
        ),
        {"user_id": alice.id},
    )
    db_session.commit()

    init_db()
    _ensure_schema_updates(get_engine())

    engine = get_engine()
    assert "requirement_file_name" in {c["name"] for c in inspect(engine).get_columns("workflows")}
    value = db_session.execute(
        text("SELECT requirement_file_name FROM workflows WHERE id = 'legacy-workflow'")
    ).scalar()
    assert value is None
